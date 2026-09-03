from __future__ import annotations

import re
import subprocess
import sys
import tkinter as tk
from datetime import date
from pathlib import Path
from tkinter import filedialog, messagebox, ttk


# ====================================================================================================
# PATHS
# ====================================================================================================

ROOT_DIR = Path(__file__).resolve().parent

DEFAULT_DAILY_TEMPLATE = (
    Path.home()
    / "Desktop"
    / "WorkSpace"
    / "DATA"
    / "26년 조기경보 매출 관리 최신.xlsx"
)

DEFAULT_DAILY_OUTPUT_DIR = (
    ROOT_DIR
    / "output"
)


# ====================================================================================================
# APP
# ====================================================================================================

class CollectorApp(ttk.Frame):

    WINDOW_WIDTH = 710
    WINDOW_HEIGHT = 420

    def __init__(
        self,
        master: tk.Tk,
    ) -> None:

        super().__init__(
            master,
            padding=10,
        )

        today = date.today()

        # ============================================================================================
        # DAILY VARIABLES
        # ============================================================================================

        self.daily_template_var = tk.StringVar(
            value=str(DEFAULT_DAILY_TEMPLATE)
        )

        self.daily_start_var = tk.StringVar(
            value=today.isoformat()
        )

        self.daily_end_var = tk.StringVar(
            value=today.isoformat()
        )

        self.daily_scope_var = tk.StringVar(
            value="all"
        )

        self.daily_store_var = tk.StringVar()

        self.daily_output_dir_var = tk.StringVar(
            value=str(DEFAULT_DAILY_OUTPUT_DIR)
        )

        # ============================================================================================
        # MENU VARIABLES
        # ============================================================================================

        self.menu_template_var = tk.StringVar()

        self.menu_year_var = tk.StringVar(
            value=str(today.year)
        )

        self.menu_month_var = tk.StringVar(
            value=str(today.month)
        )

        self.menu_scope_var = tk.StringVar(
            value="all"
        )

        self.menu_store_var = tk.StringVar()

        # Complete XLSX output path.
        self.menu_output_var = tk.StringVar()

        # Track whether the current menu output was automatically suggested.
        self._menu_output_auto_value = ""

        self._build()


    # ================================================================================================
    # WINDOW
    # ================================================================================================

    def _build(self) -> None:

        self.master.title(
            "MagicERP 매출 자동수집 W7"
        )

        self.master.geometry(
            f"{self.WINDOW_WIDTH}x{self.WINDOW_HEIGHT}"
        )

        self.master.resizable(
            False,
            False,
        )

        self.grid(
            row=0,
            column=0,
            sticky="nsew",
        )

        self.master.columnconfigure(
            0,
            weight=1,
        )

        self.master.rowconfigure(
            0,
            weight=1,
        )

        self.columnconfigure(
            0,
            weight=1,
        )

        self.rowconfigure(
            0,
            weight=1,
        )

        notebook = ttk.Notebook(
            self
        )

        notebook.grid(
            row=0,
            column=0,
            sticky="nsew",
        )

        daily_tab = ttk.Frame(
            notebook,
            padding=14,
        )

        menu_tab = ttk.Frame(
            notebook,
            padding=14,
        )

        notebook.add(
            daily_tab,
            text="매출 수집",
        )

        notebook.add(
            menu_tab,
            text="메뉴 분석",
        )

        self._build_daily_tab(
            daily_tab
        )

        self._build_menu_tab(
            menu_tab
        )


    # ================================================================================================
    # COMMON LAYOUT
    # ================================================================================================

    @staticmethod
    def _configure_tab(
        frame: ttk.Frame,
    ) -> None:

        frame.columnconfigure(
            0,
            minsize=95,
        )

        frame.columnconfigure(
            1,
            weight=1,
        )

        frame.columnconfigure(
            2,
            minsize=105,
        )


    @staticmethod
    def _label(
        frame: ttk.Frame,
        text: str,
        row: int,
    ) -> None:

        ttk.Label(
            frame,
            text=text,
        ).grid(
            row=row,
            column=0,
            sticky="w",
            pady=7,
        )


    # ================================================================================================
    # DAILY TAB
    # ================================================================================================

    def _build_daily_tab(
        self,
        frame: ttk.Frame,
    ) -> None:

        self._configure_tab(
            frame
        )

        # --------------------------------------------------------------------------------------------
        # 원본 파일
        # --------------------------------------------------------------------------------------------

        self._label(
            frame,
            "원본 파일",
            0,
        )

        ttk.Entry(
            frame,
            textvariable=self.daily_template_var,
        ).grid(
            row=0,
            column=1,
            sticky="ew",
            padx=(
                0,
                8,
            ),
            pady=7,
        )

        ttk.Button(
            frame,
            text="찾아보기",
            width=10,
            command=self._choose_daily_template,
        ).grid(
            row=0,
            column=2,
            sticky="ew",
            pady=7,
        )

        # --------------------------------------------------------------------------------------------
        # 조회 시작일
        # --------------------------------------------------------------------------------------------

        self._label(
            frame,
            "조회 시작일",
            1,
        )

        ttk.Entry(
            frame,
            textvariable=self.daily_start_var,
            width=18,
        ).grid(
            row=1,
            column=1,
            sticky="w",
            pady=7,
        )

        # --------------------------------------------------------------------------------------------
        # 조회 종료일
        # --------------------------------------------------------------------------------------------

        self._label(
            frame,
            "조회 종료일",
            2,
        )

        ttk.Entry(
            frame,
            textvariable=self.daily_end_var,
            width=18,
        ).grid(
            row=2,
            column=1,
            sticky="w",
            pady=7,
        )

        # --------------------------------------------------------------------------------------------
        # 대상 가맹점
        # --------------------------------------------------------------------------------------------

        self._label(
            frame,
            "대상 가맹점",
            3,
        )

        scope = ttk.Frame(
            frame
        )

        scope.grid(
            row=3,
            column=1,
            columnspan=2,
            sticky="w",
            pady=7,
        )

        ttk.Radiobutton(
            scope,
            text="전체 가맹점",
            variable=self.daily_scope_var,
            value="all",
            command=self._update_daily_scope,
        ).pack(
            side="left",
        )

        ttk.Radiobutton(
            scope,
            text="특정 가맹점",
            variable=self.daily_scope_var,
            value="specific",
            command=self._update_daily_scope,
        ).pack(
            side="left",
            padx=(
                25,
                0,
            ),
        )

        # --------------------------------------------------------------------------------------------
        # 매장명
        # --------------------------------------------------------------------------------------------

        self._label(
            frame,
            "매장명",
            4,
        )

        self.daily_store_entry = ttk.Entry(
            frame,
            textvariable=self.daily_store_var,
        )

        self.daily_store_entry.grid(
            row=4,
            column=1,
            columnspan=2,
            sticky="ew",
            pady=7,
        )

        # --------------------------------------------------------------------------------------------
        # 저장 위치
        # --------------------------------------------------------------------------------------------

        self._label(
            frame,
            "저장 위치",
            5,
        )

        ttk.Entry(
            frame,
            textvariable=self.daily_output_dir_var,
        ).grid(
            row=5,
            column=1,
            sticky="ew",
            padx=(
                0,
                8,
            ),
            pady=7,
        )

        ttk.Button(
            frame,
            text="찾아보기",
            width=10,
            command=self._choose_daily_output_dir,
        ).grid(
            row=5,
            column=2,
            sticky="ew",
            pady=7,
        )

        # --------------------------------------------------------------------------------------------
        # Footer
        # --------------------------------------------------------------------------------------------

        ttk.Label(
            frame,
            text="기간과 대상을 선택한 뒤 실행하세요.",
            foreground="#1976D2",
        ).grid(
            row=6,
            column=0,
            columnspan=2,
            sticky="w",
            pady=(
                24,
                0,
            ),
        )

        ttk.Button(
            frame,
            text="실행",
            width=20,
            command=self._launch_daily,
        ).grid(
            row=6,
            column=2,
            sticky="e",
            pady=(
                24,
                0,
            ),
        )

        self._update_daily_scope()


    # ================================================================================================
    # MENU TAB
    # ================================================================================================

    def _build_menu_tab(
        self,
        frame: ttk.Frame,
    ) -> None:

        self._configure_tab(
            frame
        )

        # --------------------------------------------------------------------------------------------
        # 원본 파일
        # --------------------------------------------------------------------------------------------

        self._label(
            frame,
            "원본 파일",
            0,
        )

        ttk.Entry(
            frame,
            textvariable=self.menu_template_var,
        ).grid(
            row=0,
            column=1,
            sticky="ew",
            padx=(
                0,
                8,
            ),
            pady=7,
        )

        ttk.Button(
            frame,
            text="찾아보기",
            width=10,
            command=self._choose_menu_template,
        ).grid(
            row=0,
            column=2,
            sticky="ew",
            pady=7,
        )

        # --------------------------------------------------------------------------------------------
        # 연도
        # --------------------------------------------------------------------------------------------

        self._label(
            frame,
            "연도",
            1,
        )

        self.menu_year_combo = ttk.Combobox(
            frame,
            textvariable=self.menu_year_var,
            values=[
                str(year)
                for year in range(
                    2024,
                    date.today().year + 2,
                )
            ],
            state="readonly",
            width=16,
        )

        self.menu_year_combo.grid(
            row=1,
            column=1,
            sticky="w",
            pady=7,
        )

        self.menu_year_combo.bind(
            "<<ComboboxSelected>>",
            self._menu_period_changed,
        )

        # --------------------------------------------------------------------------------------------
        # 월
        # --------------------------------------------------------------------------------------------

        self._label(
            frame,
            "월",
            2,
        )

        self.menu_month_combo = ttk.Combobox(
            frame,
            textvariable=self.menu_month_var,
            values=[
                str(month)
                for month in range(
                    1,
                    13,
                )
            ],
            state="readonly",
            width=16,
        )

        self.menu_month_combo.grid(
            row=2,
            column=1,
            sticky="w",
            pady=7,
        )

        self.menu_month_combo.bind(
            "<<ComboboxSelected>>",
            self._menu_period_changed,
        )

        # --------------------------------------------------------------------------------------------
        # 대상 가맹점
        # --------------------------------------------------------------------------------------------

        self._label(
            frame,
            "대상 가맹점",
            3,
        )

        scope = ttk.Frame(
            frame
        )

        scope.grid(
            row=3,
            column=1,
            columnspan=2,
            sticky="w",
            pady=7,
        )

        ttk.Radiobutton(
            scope,
            text="전체 가맹점",
            variable=self.menu_scope_var,
            value="all",
            command=self._update_menu_scope,
        ).pack(
            side="left",
        )

        ttk.Radiobutton(
            scope,
            text="특정 가맹점",
            variable=self.menu_scope_var,
            value="specific",
            command=self._update_menu_scope,
        ).pack(
            side="left",
            padx=(
                25,
                0,
            ),
        )

        # --------------------------------------------------------------------------------------------
        # 매장명
        # --------------------------------------------------------------------------------------------

        self._label(
            frame,
            "매장명",
            4,
        )

        self.menu_store_entry = ttk.Entry(
            frame,
            textvariable=self.menu_store_var,
        )

        self.menu_store_entry.grid(
            row=4,
            column=1,
            columnspan=2,
            sticky="ew",
            pady=7,
        )

        # --------------------------------------------------------------------------------------------
        # 저장 파일
        # --------------------------------------------------------------------------------------------

        self._label(
            frame,
            "저장 파일",
            5,
        )

        ttk.Entry(
            frame,
            textvariable=self.menu_output_var,
        ).grid(
            row=5,
            column=1,
            sticky="ew",
            padx=(
                0,
                8,
            ),
            pady=7,
        )

        ttk.Button(
            frame,
            text="찾아보기",
            width=10,
            command=self._choose_menu_output_file,
        ).grid(
            row=5,
            column=2,
            sticky="ew",
            pady=7,
        )

        # --------------------------------------------------------------------------------------------
        # Footer
        # --------------------------------------------------------------------------------------------

        self.menu_info_label = ttk.Label(
            frame,
            text=(
                "전체 가맹점 또는 특정 가맹점을 선택한 뒤 실행하세요."
            ),
            foreground="#1976D2",
        )

        self.menu_info_label.grid(
            row=6,
            column=0,
            columnspan=2,
            sticky="w",
            pady=(
                24,
                0,
            ),
        )

        ttk.Button(
            frame,
            text="실행",
            width=20,
            command=self._launch_menu,
        ).grid(
            row=6,
            column=2,
            sticky="e",
            pady=(
                24,
                0,
            ),
        )

        self._update_menu_scope()


    # ================================================================================================
    # DAILY FILE / DIRECTORY
    # ================================================================================================

    def _choose_daily_template(
        self,
    ) -> None:

        current = (
            self.daily_template_var
            .get()
            .strip()
        )

        kwargs: dict[str, str] = {
            "title": "매출 수집 원본 Excel 선택",
        }

        if current:

            current_path = Path(
                current
            )

            if current_path.parent.is_dir():

                kwargs["initialdir"] = str(
                    current_path.parent
                )

        selected = filedialog.askopenfilename(
            filetypes=[
                (
                    "Excel workbook",
                    "*.xlsx",
                ),
            ],
            **kwargs,
        )

        if selected:

            self.daily_template_var.set(
                selected
            )


    def _choose_daily_output_dir(
        self,
    ) -> None:

        current = (
            self.daily_output_dir_var
            .get()
            .strip()
        )

        kwargs: dict[str, str] = {
            "title": "매출 수집 저장 위치 선택",
        }

        if current and Path(current).is_dir():

            kwargs["initialdir"] = current

        selected = filedialog.askdirectory(
            **kwargs
        )

        if selected:

            self.daily_output_dir_var.set(
                selected
            )


    # ================================================================================================
    # MENU SOURCE
    # ================================================================================================

    def _choose_menu_template(
        self,
    ) -> None:

        current = (
            self.menu_template_var
            .get()
            .strip()
        )

        kwargs: dict[str, str] = {
            "title": "메뉴 분석 원본 Excel 선택",
        }

        if current:

            current_path = Path(
                current
            )

            if current_path.parent.is_dir():

                kwargs["initialdir"] = str(
                    current_path.parent
                )

        selected = filedialog.askopenfilename(
            filetypes=[
                (
                    "Excel workbook",
                    "*.xlsx",
                ),
            ],
            **kwargs,
        )

        if not selected:
            return

        self.menu_template_var.set(
            selected
        )

        template = Path(
            selected
        )

        # Automatically detect year/month from file name where possible.
        self._apply_period_from_filename(
            template.name
        )

        self._set_default_menu_output(
            force=True
        )


    def _apply_period_from_filename(
        self,
        filename: str,
    ) -> None:

        # Examples:
        #
        #   26년 07월 대전통합본.xlsx
        #   2026년 07월 ...
        #
        # Year is optional. Month is the most important value.

        month_match = re.search(
            r"(?<!\d)(0?[1-9]|1[0-2])월",
            filename,
        )

        if month_match:

            self.menu_month_var.set(
                str(
                    int(
                        month_match.group(1)
                    )
                )
            )

        year_4 = re.search(
            r"(?<!\d)(20\d{2})년",
            filename,
        )

        if year_4:

            self.menu_year_var.set(
                year_4.group(1)
            )

            return

        year_2 = re.search(
            r"(?<!\d)(\d{2})년",
            filename,
        )

        if year_2:

            value = int(
                year_2.group(1)
            )

            if 0 <= value <= 99:

                self.menu_year_var.set(
                    str(
                        2000 + value
                    )
                )


    def _menu_period_changed(
        self,
        _event=None,
    ) -> None:

        self._set_default_menu_output(
            force=False
        )


    # ================================================================================================
    # MENU OUTPUT FILE
    # ================================================================================================

    def _default_menu_output_path(
        self,
    ) -> Path | None:

        template_text = (
            self.menu_template_var
            .get()
            .strip()
        )

        if not template_text:

            return None

        template = Path(
            template_text
        )

        try:

            year = int(
                self.menu_year_var
                .get()
                .strip()
            )

            month = int(
                self.menu_month_var
                .get()
                .strip()
            )

        except ValueError:

            return None

        name = (
            f"{template.stem}"
            f"_메뉴분석_{year}{month:02d}.xlsx"
        )

        return (
            template.parent
            / name
        )


    def _set_default_menu_output(
        self,
        *,
        force: bool,
    ) -> None:

        proposed = (
            self._default_menu_output_path()
        )

        if proposed is None:
            return

        current = (
            self.menu_output_var
            .get()
            .strip()
        )

        if (
            force
            or not current
            or current == self._menu_output_auto_value
        ):

            value = str(
                proposed
            )

            self.menu_output_var.set(
                value
            )

            self._menu_output_auto_value = (
                value
            )


    def _choose_menu_output_file(
        self,
    ) -> None:

        proposed = (
            self._default_menu_output_path()
        )

        current_text = (
            self.menu_output_var
            .get()
            .strip()
        )

        if current_text:

            current = Path(
                current_text
            )

            initial_dir = (
                current.parent
                if current.parent.is_dir()
                else Path.home() / "Desktop"
            )

            initial_file = (
                current.name
                if current.suffix.lower() == ".xlsx"
                else "메뉴분석.xlsx"
            )

        elif proposed is not None:

            initial_dir = (
                proposed.parent
            )

            initial_file = (
                proposed.name
            )

        else:

            initial_dir = (
                Path.home()
                / "Desktop"
            )

            initial_file = (
                "메뉴분석.xlsx"
            )

        selected = filedialog.asksaveasfilename(
            title="메뉴 분석 결과 저장",
            initialdir=str(
                initial_dir
            ),
            initialfile=initial_file,
            defaultextension=".xlsx",
            filetypes=[
                (
                    "Excel workbook",
                    "*.xlsx",
                ),
            ],
        )

        if not selected:
            return

        output = Path(
            selected
        )

        if output.suffix.lower() != ".xlsx":

            output = output.with_suffix(
                ".xlsx"
            )

        self.menu_output_var.set(
            str(output)
        )

        # User explicitly selected this value.
        self._menu_output_auto_value = ""


    # ================================================================================================
    # SCOPE STATE
    # ================================================================================================

    def _update_daily_scope(
        self,
    ) -> None:

        if self.daily_scope_var.get() == "specific":

            self.daily_store_entry.configure(
                state="normal"
            )

        else:

            self.daily_store_entry.configure(
                state="disabled"
            )


    def _update_menu_scope(
        self,
    ) -> None:

        if self.menu_scope_var.get() == "specific":

            self.menu_store_entry.configure(
                state="normal"
            )

            self.menu_info_label.configure(
                text=(
                    "특정 가맹점은 쉼표로 구분하세요. "
                    "예: 세천점, 파호점"
                )
            )

        else:

            self.menu_store_entry.configure(
                state="disabled"
            )

            self.menu_info_label.configure(
                text=(
                    "원본 Excel의 전체 가맹점을 대상으로 실행합니다."
                )
            )


    # ================================================================================================
    # COMMAND TARGET
    # ================================================================================================

    @staticmethod
    def _collector_command_base(
    ) -> list[str]:

        # --------------------------------------------------------------------------------------------
        # Development:
        #
        #   python gui.py
        #
        # launches:
        #
        #   python main.py ...
        #
        # --------------------------------------------------------------------------------------------

        if not getattr(
            sys,
            "frozen",
            False,
        ):

            return [
                sys.executable,
                str(
                    Path(__file__)
                    .resolve()
                    .with_name("main.py")
                ),
            ]

        # --------------------------------------------------------------------------------------------
        # Frozen GUI:
        #
        # release\
        #   MagicERP_AutoCollector.exe
        #   MagicERP_AutoCollector_GUI.exe
        #
        # GUI launches the console collector EXE.
        # --------------------------------------------------------------------------------------------

        exe_dir = (
            Path(
                sys.executable
            )
            .resolve()
            .parent
        )

        candidates = [
            exe_dir
            / "MagicERP_AutoCollector.exe",

            exe_dir
            / "Sales_Data_Collector.exe",

            exe_dir
            / "Sales_Data_Collector",
        ]

        for candidate in candidates:

            if candidate.is_file():

                return [
                    str(
                        candidate
                    )
                ]

        raise FileNotFoundError(
            (
                "수집 실행파일을 찾을 수 없습니다.\n\n"
                f"GUI 위치: {exe_dir}\n\n"
                "MagicERP_AutoCollector.exe가 "
                "GUI와 같은 폴더에 있어야 합니다."
            )
        )


    # ================================================================================================
    # CONSOLE LAUNCH
    # ================================================================================================

    @staticmethod
    def _launch_collector_console(
        command: list[str],
    ) -> None:

        # IMPORTANT:
        #
        # Do NOT use:
        #
        #   Do not route the command through cmd.exe string conversion.
        #
        # That makes Windows parse paths with spaces and Korean text
        # for a second time.
        #
        # Passing a list directly lets subprocess handle argument quoting.
        #
        # CREATE_NEW_CONSOLE creates an interactive console so getpass()
        # and input() work normally:
        #
        #   Service ID:
        #   Password:
        #
        # Login succeeds -> collector continues automatically.

        try:

            creation_flags = getattr(
                subprocess,
                "CREATE_NEW_CONSOLE",
                0,
            )

            subprocess.Popen(
                command,
                creationflags=creation_flags,
                cwd=str(
                    ROOT_DIR
                ),
            )

        except Exception as exc:

            messagebox.showerror(
                "실행 오류",
                str(exc),
            )


    # ================================================================================================
    # DAILY EXECUTION
    # ================================================================================================

    def _launch_daily(
        self,
    ) -> None:

        template_text = (
            self.daily_template_var
            .get()
            .strip()
        )

        template = Path(
            template_text
        )

        if not template.is_file():

            messagebox.showerror(
                "원본 파일",
                "유효한 Excel 원본 파일을 선택하세요.",
            )

            return

        try:

            start = date.fromisoformat(
                self.daily_start_var
                .get()
                .strip()
            )

            end = date.fromisoformat(
                self.daily_end_var
                .get()
                .strip()
            )

        except ValueError:

            messagebox.showerror(
                "기간 확인",
                "날짜는 YYYY-MM-DD 형식으로 입력하세요.",
            )

            return

        if start > end:

            messagebox.showerror(
                "기간 확인",
                "조회 시작일이 조회 종료일보다 늦습니다.",
            )

            return

        output_dir_text = (
            self.daily_output_dir_var
            .get()
            .strip()
        )

        if not output_dir_text:

            messagebox.showerror(
                "저장 위치",
                "저장 위치를 선택하세요.",
            )

            return

        output_dir = Path(
            output_dir_text
        )

        try:

            output_dir.mkdir(
                parents=True,
                exist_ok=True,
            )

        except OSError as exc:

            messagebox.showerror(
                "저장 위치",
                str(exc),
            )

            return

        try:

            command = (
                self._collector_command_base()
            )

        except Exception as exc:

            messagebox.showerror(
                "실행 파일",
                str(exc),
            )

            return

        command.extend(
            [
                "--production-write",
                "--template",
                str(template),
                "--start-date",
                start.isoformat(),
                "--end-date",
                end.isoformat(),
                "--production-output",
                str(output_dir),
            ]
        )

        if self.daily_scope_var.get() == "all":

            command.append(
                "--all-stores"
            )

        else:

            store_name = (
                self.daily_store_var
                .get()
                .strip()
            )

            if not store_name:

                messagebox.showerror(
                    "가맹점 확인",
                    "가맹점명을 입력하세요.",
                )

                return

            command.extend(
                [
                    "--store-name",
                    store_name,
                ]
            )

        self._launch_collector_console(
            command
        )


    # ================================================================================================
    # MENU EXECUTION
    # ================================================================================================

    def _launch_menu(
        self,
    ) -> None:

        # --------------------------------------------------------------------------------------------
        # Source workbook
        # --------------------------------------------------------------------------------------------

        template_text = (
            self.menu_template_var
            .get()
            .strip()
        )

        template = Path(
            template_text
        )

        if not template.is_file():

            messagebox.showerror(
                "원본 파일",
                "유효한 메뉴 통합 Excel 파일을 선택하세요.",
            )

            return

        # --------------------------------------------------------------------------------------------
        # Period
        # --------------------------------------------------------------------------------------------

        try:

            year = int(
                self.menu_year_var
                .get()
                .strip()
            )

            month = int(
                self.menu_month_var
                .get()
                .strip()
            )

        except ValueError:

            messagebox.showerror(
                "기간 확인",
                "연도와 월을 확인하세요.",
            )

            return

        if not 1 <= month <= 12:

            messagebox.showerror(
                "기간 확인",
                "월은 1~12 사이여야 합니다.",
            )

            return

        # --------------------------------------------------------------------------------------------
        # Output file
        # --------------------------------------------------------------------------------------------

        output_text = (
            self.menu_output_var
            .get()
            .strip()
        )

        if not output_text:

            self._set_default_menu_output(
                force=True
            )

            output_text = (
                self.menu_output_var
                .get()
                .strip()
            )

        if not output_text:

            messagebox.showerror(
                "저장 파일",
                "저장할 Excel 파일명과 위치를 지정하세요.",
            )

            return

        output = Path(
            output_text
        )

        if output.suffix.lower() != ".xlsx":

            output = output.with_suffix(
                ".xlsx"
            )

            self.menu_output_var.set(
                str(output)
            )

        # --------------------------------------------------------------------------------------------
        # Original protection
        # --------------------------------------------------------------------------------------------

        try:

            if (
                template.resolve()
                == output.resolve()
            ):

                messagebox.showerror(
                    "저장 파일",
                    "원본 Excel 파일과 동일한 경로로 저장할 수 없습니다.",
                )

                return

        except OSError:
            pass

        # --------------------------------------------------------------------------------------------
        # Parent directory
        # --------------------------------------------------------------------------------------------

        try:

            output.parent.mkdir(
                parents=True,
                exist_ok=True,
            )

        except OSError as exc:

            messagebox.showerror(
                "저장 파일",
                (
                    "저장 폴더를 생성할 수 없습니다.\n\n"
                    f"{exc}"
                ),
            )

            return

        # --------------------------------------------------------------------------------------------
        # Existing output
        #
        # menu_monthly_batch.py protects existing outputs.
        # Ask user first, then remove only the generated target.
        # --------------------------------------------------------------------------------------------

        if output.exists():

            overwrite = messagebox.askyesno(
                "저장 파일",
                (
                    "동일한 결과 파일이 이미 존재합니다.\n\n"
                    f"{output}\n\n"
                    "기존 결과 파일을 삭제하고 새로 생성하시겠습니까?"
                ),
            )

            if not overwrite:
                return

            try:

                output.unlink()

            except OSError as exc:

                messagebox.showerror(
                    "저장 파일",
                    (
                        "기존 결과 파일을 삭제할 수 없습니다.\n"
                        "Excel에서 파일이 열려 있는지 확인하세요.\n\n"
                        f"{exc}"
                    ),
                )

                return

        # --------------------------------------------------------------------------------------------
        # Store scope
        # --------------------------------------------------------------------------------------------

        if self.menu_scope_var.get() == "all":

            store_names = (
                self._read_workbook_store_names(
                    template
                )
            )

            if not store_names:

                messagebox.showerror(
                    "가맹점 확인",
                    "원본 Excel에서 처리할 가맹점을 찾지 못했습니다.",
                )

                return

        else:

            raw = (
                self.menu_store_var
                .get()
                .strip()
            )

            store_names = []

            seen: set[str] = set()

            for value in raw.split(","):

                name = value.strip()

                if not name:
                    continue

                if name in seen:
                    continue

                seen.add(
                    name
                )

                store_names.append(
                    name
                )

            if not store_names:

                messagebox.showerror(
                    "가맹점 확인",
                    (
                        "한 개 이상의 가맹점명을 입력하세요.\n\n"
                        "여러 매장은 쉼표로 구분합니다.\n"
                        "예: 세천점, 파호점"
                    ),
                )

                return

        # --------------------------------------------------------------------------------------------
        # Confirmation
        # --------------------------------------------------------------------------------------------

        confirmed = messagebox.askyesno(
            "메뉴 분석 실행",
            (
                f"조회 기간 : {year}-{month:02d}\n"
                f"대상 매장 : {len(store_names)}개\n\n"
                f"저장 파일:\n{output}\n\n"
                "실행하면 새 CMD 창에서 "
                "Service ID와 Password를 입력합니다.\n\n"
                "계속하시겠습니까?"
            ),
        )

        if not confirmed:
            return

        # --------------------------------------------------------------------------------------------
        # Command
        # --------------------------------------------------------------------------------------------

        try:

            command = (
                self._collector_command_base()
            )

        except Exception as exc:

            messagebox.showerror(
                "실행 파일",
                str(exc),
            )

            return

        command.extend(
            [
                "--menu-monthly-batch",
                "--template",
                str(template),
                "--output",
                str(output),
                "--year",
                str(year),
                "--month",
                str(month),
            ]
        )

        for store_name in store_names:

            command.extend(
                [
                    "--store-name",
                    store_name,
                ]
            )

        self._launch_collector_console(
            command
        )


    # ================================================================================================
    # MENU STORE DISCOVERY
    # ================================================================================================

    @staticmethod
    def _normalize_cell(
        value,
    ) -> str:

        if value is None:
            return ""

        return (
            str(value)
            .replace(
                "\u00a0",
                " ",
            )
            .replace(
                "\r",
                "",
            )
            .replace(
                "\n",
                "",
            )
            .strip()
        )


    @classmethod
    def _read_workbook_store_names(
        cls,
        template: Path,
    ) -> list[str]:

        try:

            from openpyxl import load_workbook

            workbook = load_workbook(
                template,
                read_only=True,
                data_only=False,
            )

        except Exception as exc:

            messagebox.showerror(
                "Excel 확인",
                (
                    "원본 Excel을 열 수 없습니다.\n\n"
                    f"{exc}"
                ),
            )

            return []

        stores: list[str] = []

        seen: set[str] = set()

        discovered: list[
            tuple[
                str,
                int,
                str,
            ]
        ] = []

        try:

            # Inspect all worksheets.
            #
            # Daejeon:
            #   07월대전통합본
            #   폼2
            #
            # Daegu:
            #   7월 통합본
            #
            # A store sales block is identified by:
            #
            #   C<row> = store name
            #   F<row> = 매출

            for worksheet in workbook.worksheets:

                for row in range(
                    1,
                    worksheet.max_row + 1,
                ):

                    marker = cls._normalize_cell(
                        worksheet[
                            f"F{row}"
                        ].value
                    )

                    if marker != "매출":
                        continue

                    store_name = cls._normalize_cell(
                        worksheet[
                            f"C{row}"
                        ].value
                    )

                    if not store_name:
                        continue

                    if store_name in seen:
                        continue

                    seen.add(
                        store_name
                    )

                    stores.append(
                        store_name
                    )

                    discovered.append(
                        (
                            worksheet.title,
                            row,
                            store_name,
                        )
                    )

        finally:

            workbook.close()

        # This output appears in the GUI parent console when developing
        # and is useful when validating regional workbook structure.

        print("")
        print(
            "=" * 100
        )
        print(
            "MENU WORKBOOK STORE DISCOVERY"
        )
        print(
            "=" * 100
        )
        print(
            f"WORKBOOK={template}"
        )
        print(
            f"STORE_COUNT={len(stores)}"
        )

        for index, (
            sheet_name,
            row,
            store_name,
        ) in enumerate(
            discovered,
            start=1,
        ):

            print(
                f"STORE_{index}="
                f"{store_name}"
                f" | SHEET={sheet_name}"
                f" | SALES_ROW={row}"
            )

        print(
            "=" * 100
        )

        return stores


# ====================================================================================================
# MAIN
# ====================================================================================================

def main() -> None:

    root = tk.Tk()

    CollectorApp(
        root
    )

    root.mainloop()


if __name__ == "__main__":
    main()