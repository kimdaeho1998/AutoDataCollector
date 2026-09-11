from __future__ import annotations

import hashlib
import shutil
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from .product_workbook_dry_run import (
    ProductWorkbookDryRunStatus,
    dry_run_product_workbook_plan,
)
from .product_workbook_plan import ProductWorkbookPlan


@dataclass(frozen=True)
class ProductWorkbookCopyWriteResult:
    source_path: Path
    output_path: Path
    store_name: str
    written_cells: int
    same_value_cells: int
    source_hash_before: str
    source_hash_after: str


class ProductWorkbookCopyWriter:
    """
    Safe copy-only writer for Product Workbook.

    Contract:
    - source workbook is never saved directly
    - output is created with shutil.copy2()
    - dry-run must allow writing
    - only READY Product targets are written
    - SAME_VALUE cells are preserved untouched
    - Z residual formulas are not written
    - ratio row formulas are not written
    - output is reopened and verified after save
    - original source hash must remain unchanged
    """

    def __init__(
        self,
        source_path: str | Path,
        *,
        sheet_name: str = "7월",
    ) -> None:
        self.source_path = Path(source_path)
        self.sheet_name = sheet_name

    def write_copy(
        self,
        *,
        plan: ProductWorkbookPlan,
        output_path: str | Path,
    ) -> ProductWorkbookCopyWriteResult:

        output = Path(output_path)

        if not self.source_path.exists():
            raise FileNotFoundError(
                f"SOURCE_WORKBOOK_NOT_FOUND:{self.source_path}"
            )

        if output.resolve() == self.source_path.resolve():
            raise ValueError(
                "OUTPUT_MUST_DIFFER_FROM_SOURCE"
            )

        source_hash_before = self._hash(
            self.source_path
        )

        source_book = load_workbook(
            self.source_path,
            data_only=False,
        )

        try:
            if self.sheet_name not in source_book.sheetnames:
                raise ValueError(
                    f"PRODUCT_SHEET_NOT_FOUND:{self.sheet_name}"
                )

            source_sheet = source_book[
                self.sheet_name
            ]

            dry_run = dry_run_product_workbook_plan(
                worksheet=source_sheet,
                plan=plan,
            )

            if not dry_run.can_write:
                raise ValueError(
                    "PRODUCT_WORKBOOK_WRITE_BLOCKED:"
                    f"CONFLICT={dry_run.conflict_count},"
                    f"BLOCKED={dry_run.blocked_count}"
                )

            source_sheetnames = tuple(
                source_book.sheetnames
            )

            source_merges = tuple(
                sorted(
                    str(item)
                    for item in source_sheet.merged_cells.ranges
                )
            )

            protected_formulas = self._protected_formula_snapshot(
                source_sheet,
                plan,
            )

        finally:
            source_book.close()

        output.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        shutil.copy2(
            self.source_path,
            output,
        )

        written = 0
        same = 0

        try:
            workbook = load_workbook(
                output,
                data_only=False,
            )

            try:
                if self.sheet_name not in workbook.sheetnames:
                    raise ValueError(
                        f"PRODUCT_SHEET_NOT_FOUND:{self.sheet_name}"
                    )

                worksheet = workbook[
                    self.sheet_name
                ]

                fresh_dry_run = dry_run_product_workbook_plan(
                    worksheet=worksheet,
                    plan=plan,
                )

                if not fresh_dry_run.can_write:
                    raise ValueError(
                        "PRODUCT_WORKBOOK_FRESH_WRITE_BLOCKED:"
                        f"CONFLICT={fresh_dry_run.conflict_count},"
                        f"BLOCKED={fresh_dry_run.blocked_count}"
                    )

                dry_by_coordinate = {
                    item.coordinate: item
                    for item in fresh_dry_run.cells
                }

                # --------------------------------------------
                # Sparse canonical E:Y
                # --------------------------------------------

                for cell in plan.cells:

                    sales_dry = dry_by_coordinate[
                        cell.sales_cell
                    ]

                    if (
                        sales_dry.status
                        == ProductWorkbookDryRunStatus.READY
                    ):
                        worksheet[
                            cell.sales_cell
                        ].value = int(
                            cell.proposed_sales
                        )
                        written += 1

                    elif (
                        sales_dry.status
                        == ProductWorkbookDryRunStatus.SAME_VALUE
                    ):
                        same += 1

                    else:
                        raise ValueError(
                            "UNEXPECTED_SALES_WRITE_STATUS:"
                            f"{cell.sales_cell}:"
                            f"{sales_dry.status.value}"
                        )

                    quantity_dry = dry_by_coordinate[
                        cell.quantity_cell
                    ]

                    if (
                        quantity_dry.status
                        == ProductWorkbookDryRunStatus.READY
                    ):
                        if cell.proposed_quantity is None:
                            raise ValueError(
                                "PRODUCT_QUANTITY_MISSING:"
                                f"{cell.quantity_cell}"
                            )

                        worksheet[
                            cell.quantity_cell
                        ].value = int(
                            cell.proposed_quantity
                        )
                        written += 1

                    elif (
                        quantity_dry.status
                        == ProductWorkbookDryRunStatus.SAME_VALUE
                    ):
                        same += 1

                    else:
                        raise ValueError(
                            "UNEXPECTED_QUANTITY_WRITE_STATUS:"
                            f"{cell.quantity_cell}:"
                            f"{quantity_dry.status.value}"
                        )

                # --------------------------------------------
                # AA totals
                # --------------------------------------------

                sales_total_dry = dry_by_coordinate[
                    plan.sales_total_cell
                ]

                if (
                    sales_total_dry.status
                    == ProductWorkbookDryRunStatus.READY
                ):
                    worksheet[
                        plan.sales_total_cell
                    ].value = int(
                        plan.source_total_sales
                    )
                    written += 1

                elif (
                    sales_total_dry.status
                    == ProductWorkbookDryRunStatus.SAME_VALUE
                ):
                    same += 1

                else:
                    raise ValueError(
                        "UNEXPECTED_SALES_TOTAL_STATUS:"
                        f"{plan.sales_total_cell}:"
                        f"{sales_total_dry.status.value}"
                    )

                quantity_total_dry = dry_by_coordinate[
                    plan.quantity_total_cell
                ]

                if (
                    quantity_total_dry.status
                    == ProductWorkbookDryRunStatus.READY
                ):
                    worksheet[
                        plan.quantity_total_cell
                    ].value = int(
                        plan.source_total_quantity
                    )
                    written += 1

                elif (
                    quantity_total_dry.status
                    == ProductWorkbookDryRunStatus.SAME_VALUE
                ):
                    same += 1

                else:
                    raise ValueError(
                        "UNEXPECTED_QUANTITY_TOTAL_STATUS:"
                        f"{plan.quantity_total_cell}:"
                        f"{quantity_total_dry.status.value}"
                    )

                workbook.save(
                    output
                )

            finally:
                workbook.close()

            verified = load_workbook(
                output,
                data_only=False,
            )

            try:
                self._verify_output(
                    workbook=verified,
                    plan=plan,
                    expected_sheetnames=source_sheetnames,
                    expected_merges=source_merges,
                    expected_protected_formulas=protected_formulas,
                )
            finally:
                verified.close()

            source_hash_after = self._hash(
                self.source_path
            )

            if source_hash_before != source_hash_after:
                raise ValueError(
                    "ORIGINAL_MODIFIED"
                )

            return ProductWorkbookCopyWriteResult(
                source_path=self.source_path,
                output_path=output,
                store_name=plan.store_name,
                written_cells=written,
                same_value_cells=same,
                source_hash_before=source_hash_before,
                source_hash_after=source_hash_after,
            )

        except Exception:
            if output.exists():
                output.unlink()
            raise

    def _verify_output(
        self,
        *,
        workbook,
        plan: ProductWorkbookPlan,
        expected_sheetnames: tuple[str, ...],
        expected_merges: tuple[str, ...],
        expected_protected_formulas: dict[str, object],
    ) -> None:

        if tuple(workbook.sheetnames) != expected_sheetnames:
            raise ValueError(
                "SHEET_ORDER_CHANGED"
            )

        worksheet = workbook[
            self.sheet_name
        ]

        current_merges = tuple(
            sorted(
                str(item)
                for item in worksheet.merged_cells.ranges
            )
        )

        if current_merges != expected_merges:
            raise ValueError(
                "MERGED_RANGES_CHANGED"
            )

        for cell in plan.cells:

            if worksheet[
                cell.sales_cell
            ].value != int(
                cell.proposed_sales
            ):
                raise ValueError(
                    "SALES_VERIFY_FAILED:"
                    f"{cell.sales_cell}"
                )

            if cell.proposed_quantity is None:
                raise ValueError(
                    "PRODUCT_QUANTITY_MISSING:"
                    f"{cell.quantity_cell}"
                )

            if worksheet[
                cell.quantity_cell
            ].value != int(
                cell.proposed_quantity
            ):
                raise ValueError(
                    "QUANTITY_VERIFY_FAILED:"
                    f"{cell.quantity_cell}"
                )

        if worksheet[
            plan.sales_total_cell
        ].value != int(
            plan.source_total_sales
        ):
            raise ValueError(
                "SALES_TOTAL_VERIFY_FAILED:"
                f"{plan.sales_total_cell}"
            )

        if worksheet[
            plan.quantity_total_cell
        ].value != int(
            plan.source_total_quantity
        ):
            raise ValueError(
                "QUANTITY_TOTAL_VERIFY_FAILED:"
                f"{plan.quantity_total_cell}"
            )

        protected_after = self._protected_formula_snapshot(
            worksheet,
            plan,
        )

        if protected_after != expected_protected_formulas:
            raise ValueError(
                "PROTECTED_FORMULA_CHANGED"
            )

    @staticmethod
    def _protected_formula_snapshot(
        worksheet,
        plan: ProductWorkbookPlan,
    ) -> dict[str, object]:

        coordinates: list[str] = [
            f"Z{plan.sales_row}",
            f"Z{plan.quantity_row}",
        ]

        # Preserve entire ratio row E:Y + Z.
        for column in range(
            5,
            27,
        ):
            coordinates.append(
                worksheet.cell(
                    row=plan.ratio_row,
                    column=column,
                ).coordinate
            )

        return {
            coordinate: worksheet[
                coordinate
            ].value
            for coordinate in coordinates
        }

    @staticmethod
    def _hash(
        path: Path,
    ) -> str:

        digest = hashlib.sha256()

        with path.open(
            "rb"
        ) as handle:

            for chunk in iter(
                lambda: handle.read(
                    1024 * 1024
                ),
                b"",
            ):
                digest.update(
                    chunk
                )

        return digest.hexdigest()
