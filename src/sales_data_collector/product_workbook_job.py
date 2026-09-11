from __future__ import annotations

import hashlib
import shutil
import time

from collections import Counter
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Sequence

from openpyxl import load_workbook

from sales_data_collector.client import ServiceClient

from sales_data_collector.excel.product_projection import (
    project_product_detail,
)

from sales_data_collector.excel.product_workbook_plan import (
    build_product_workbook_plan,
    normalize_product_store_name,
)

from sales_data_collector.product_workbook_batch import (
    read_product_workbook_store_targets,
)


@dataclass(frozen=True)
class ProductWorkbookJobResult:
    source_path: Path
    output_path: Path

    target_count: int
    written_count: int
    empty_count: int
    erp_not_found_count: int
    erp_ambiguous_count: int
    error_count: int

    written_cell_count: int

    elapsed_seconds: float

    source_sha256: str
    output_sha256: str

    failures: tuple[
        tuple[str, str],
        ...
    ]


def _sha256(
    path: Path,
) -> str:

    digest = hashlib.sha256()

    with path.open(
        "rb"
    ) as handle:

        for chunk in iter(
            lambda: handle.read(
                1024 * 1024
            ),
            b"",
        ):
            digest.update(
                chunk
            )

    return (
        digest
        .hexdigest()
        .upper()
    )


def _is_formula(
    value: object,
) -> bool:

    return (
        isinstance(
            value,
            str,
        )
        and value.startswith(
            "="
        )
    )


def _has_product_data(
    result,
) -> bool:
    """
    Product existence is based on parsed row presence.

    Zero and negative sales / quantity are valid source values.
    """

    return bool(
        result.records
    )


def _collect_product_detail(
    client: ServiceClient,
    *,
    brand_idx: str,
    brand_name: str,
    store,
    start_date: date,
    end_date: date,
):

    return (
        client
        .get_product_detail_sales(
            start_date=start_date,
            end_date=end_date,
            brand_idx=brand_idx,
            brand_name=brand_name,
            store_idx=(
                store.magic_store_id
            ),
            store_name=(
                store.store_name
            ),
        )
    )


def _validate_plan_before_write(
    worksheet,
    plan,
) -> None:
    """
    Validate every target before mutating a store block.

    This prevents partial per-store writes.
    """

    for cell_plan in (
        plan.cells
    ):

        if (
            cell_plan
            .proposed_quantity
            is None
        ):
            raise RuntimeError(
                "PRODUCT_QUANTITY_MISSING:"
                f"{plan.store_name}:"
                f"{cell_plan.canonical_key}"
            )

        for address in (
            cell_plan.sales_cell,
            cell_plan.quantity_cell,
        ):

            if _is_formula(
                worksheet[
                    address
                ].value
            ):
                raise RuntimeError(
                    "DIRECT_TARGET_IS_FORMULA:"
                    f"{plan.store_name}:"
                    f"{address}"
                )

    if (
        plan.source_total_quantity
        is None
    ):
        raise RuntimeError(
            "PRODUCT_TOTAL_QUANTITY_MISSING:"
            f"{plan.store_name}"
        )

    for address in (
        plan.sales_total_cell,
        plan.quantity_total_cell,
    ):

        if _is_formula(
            worksheet[
                address
            ].value
        ):
            raise RuntimeError(
                "TOTAL_TARGET_IS_FORMULA:"
                f"{plan.store_name}:"
                f"{address}"
            )


def _apply_plan(
    worksheet,
    plan,
) -> int:
    """
    Write only Product Workbook authoritative value targets.

    E:Y
        canonical product sales / quantity

    AA
        workbook-eligible total sales / quantity

    Z
        untouched
        existing AA-SUM(E:Y) formula remains authoritative

    ratio row
        untouched
    """

    _validate_plan_before_write(
        worksheet,
        plan,
    )

    written = 0

    for cell_plan in (
        plan.cells
    ):

        worksheet[
            cell_plan.sales_cell
        ] = int(
            cell_plan.proposed_sales
        )

        worksheet[
            cell_plan.quantity_cell
        ] = int(
            cell_plan.proposed_quantity
        )

        written += 2

    worksheet[
        plan.sales_total_cell
    ] = int(
        plan.source_total_sales
    )

    worksheet[
        plan.quantity_total_cell
    ] = int(
        plan.source_total_quantity
    )

    written += 2

    return written


def _resolve_seojeongri(
    client: ServiceClient,
    *,
    brand_idx: str,
    brand_name: str,
    candidates: Sequence,
    start_date: date,
    end_date: date,
):
    """
    Resolve the confirmed duplicated ERP store name
    서정리역점 using parsed Product Detail row presence.
    """

    with_data = []

    for candidate in (
        candidates
    ):

        result = (
            _collect_product_detail(
                client,
                brand_idx=brand_idx,
                brand_name=brand_name,
                store=candidate,
                start_date=start_date,
                end_date=end_date,
            )
        )

        if _has_product_data(
            result
        ):
            with_data.append(
                (
                    candidate,
                    result,
                )
            )

    if len(
        with_data
    ) == 1:
        return with_data[0]

    if not with_data:
        return None

    raise RuntimeError(
        "SEOJEONGRI_MULTIPLE_DATA_CANDIDATES:"
        + ",".join(
            store.magic_store_id
            for store, _
            in with_data
        )
    )


def run_product_workbook_job(
    *,
    client: ServiceClient,
    erp_stores: Sequence,
    brand_idx: str,
    brand_name: str,
    source_path: str | Path,
    output_path: str | Path,
    start_date: date,
    end_date: date,
    sheet_name: str,
    progress: bool = True,
) -> ProductWorkbookJobResult:
    """
    Collect Product Detail sequentially for every store represented
    in the Product Workbook and write a new workbook copy.

    Contract
    --------
    - source workbook is never modified
    - workbook is copied once
    - output workbook is opened once
    - Product Detail is collected sequentially
    - only successfully collected stores are updated
    - EMPTY / ERP_NOT_FOUND stores preserve existing workbook values
    - unknown duplicate stores fail closed
    - per-store errors preserve that store's original workbook values
    - Z formulas and ratio formulas are preserved
    - negative source values are preserved
    - output workbook is saved once
    """

    started = (
        time.perf_counter()
    )

    source = Path(
        source_path
    )

    output = Path(
        output_path
    )

    if not source.exists():
        raise FileNotFoundError(
            f"SOURCE_NOT_FOUND:{source}"
        )

    if (
        source.resolve()
        == output.resolve()
    ):
        raise RuntimeError(
            "OUTPUT_MUST_DIFFER_FROM_SOURCE"
        )

    if (
        end_date
        < start_date
    ):
        raise ValueError(
            "INVALID_PRODUCT_PERIOD"
        )

    source_hash_before = (
        _sha256(
            source
        )
    )

    # ========================================================
    # ERP STORE INDEX
    # ========================================================

    erp_by_name: dict[
        str,
        list,
    ] = {}

    for store in (
        erp_stores
    ):

        normalized = (
            normalize_product_store_name(
                store.store_name
            )
        )

        if not normalized:
            continue

        erp_by_name.setdefault(
            normalized,
            [],
        ).append(
            store
        )

    # ========================================================
    # ONE OUTPUT COPY
    # ========================================================

    output.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    if output.exists():
        output.unlink()

    shutil.copy2(
        source,
        output,
    )

    workbook = load_workbook(
        output,
        data_only=False,
    )

    status = Counter()

    failures: list[
        tuple[str, str]
    ] = []

    written_cell_count = 0

    try:

        if (
            sheet_name
            not in workbook.sheetnames
        ):
            raise RuntimeError(
                "PRODUCT_WORKBOOK_SHEET_NOT_FOUND:"
                f"{sheet_name}"
            )

        worksheet = workbook[
            sheet_name
        ]

        targets = (
            read_product_workbook_store_targets(
                worksheet
            )
        )

        total = len(
            targets
        )

        # ====================================================
        # STORE LOOP
        # ====================================================

        for index, target in enumerate(
            targets,
            start=1,
        ):

            store_name = (
                target.store_name
            )

            normalized = (
                normalize_product_store_name(
                    store_name
                )
            )

            if progress:
                print(
                    f"[{index:03d}/{total:03d}] "
                    f"{store_name}"
                )

            candidates = (
                erp_by_name.get(
                    normalized,
                    [],
                )
            )

            if not candidates:

                status[
                    "ERP_NOT_FOUND"
                ] += 1

                if progress:
                    print(
                        "  -> ERP_NOT_FOUND"
                    )

                continue

            try:

                # ============================================
                # UNIQUE STORE
                # ============================================

                if len(
                    candidates
                ) == 1:

                    store = (
                        candidates[0]
                    )

                    result = (
                        _collect_product_detail(
                            client,
                            brand_idx=brand_idx,
                            brand_name=brand_name,
                            store=store,
                            start_date=start_date,
                            end_date=end_date,
                        )
                    )

                # ============================================
                # CONFIRMED SEOJEONGRI DUPLICATE
                # ============================================

                elif (
                    "".join(
                        normalized.split()
                    )
                    == "서정리역점"
                ):

                    resolved = (
                        _resolve_seojeongri(
                            client,
                            brand_idx=brand_idx,
                            brand_name=brand_name,
                            candidates=candidates,
                            start_date=start_date,
                            end_date=end_date,
                        )
                    )

                    if resolved is None:

                        status[
                            "EMPTY"
                        ] += 1

                        if progress:
                            print(
                                "  -> EMPTY"
                            )

                        continue

                    (
                        store,
                        result,
                    ) = resolved

                # ============================================
                # UNKNOWN DUPLICATE
                # ============================================

                else:

                    ids = ",".join(
                        candidate
                        .magic_store_id
                        for candidate
                        in candidates
                    )

                    reason = (
                        "ERP_AMBIGUOUS:"
                        + ids
                    )

                    status[
                        "ERP_AMBIGUOUS"
                    ] += 1

                    failures.append(
                        (
                            store_name,
                            reason,
                        )
                    )

                    if progress:
                        print(
                            f"  -> {reason}"
                        )

                    continue

                # ============================================
                # VALID EMPTY
                # ============================================

                if not (
                    _has_product_data(
                        result
                    )
                ):

                    status[
                        "EMPTY"
                    ] += 1

                    if progress:
                        print(
                            "  -> EMPTY"
                        )

                    continue

                # ============================================
                # RAW -> PROJECTION
                # ============================================

                projection = (
                    project_product_detail(
                        result.records
                    )
                )

                # ============================================
                # PROJECTION -> PLAN
                # ============================================

                plan = (
                    build_product_workbook_plan(
                        worksheet=worksheet,
                        store_name=store_name,
                        projection=projection,
                        source_total_sales=(
                            result
                            .source_total_sales
                        ),
                        source_total_quantity=(
                            result
                            .source_total_quantity
                        ),
                    )
                )

                # ============================================
                # PLAN -> WORKBOOK MEMORY
                # ============================================

                written = (
                    _apply_plan(
                        worksheet,
                        plan,
                    )
                )

                written_cell_count += (
                    written
                )

                status[
                    "WRITTEN"
                ] += 1

                if progress:

                    print(
                        "  -> WRITTEN "
                        f"records="
                        f"{len(result.records)} "
                        f"mapped="
                        f"{len(projection.mapped)} "
                        f"unmapped="
                        f"{len(projection.unmapped)} "
                        f"AA_sales="
                        f"{plan.source_total_sales} "
                        f"AA_qty="
                        f"{plan.source_total_quantity}"
                    )

            except Exception as exc:

                reason = (
                    f"{type(exc).__name__}:"
                    f"{exc}"
                )

                status[
                    "ERROR"
                ] += 1

                failures.append(
                    (
                        store_name,
                        reason,
                    )
                )

                if progress:
                    print(
                        f"  -> ERROR "
                        f"{reason}"
                    )

                upper = (
                    str(exc)
                    .upper()
                )

                # Fatal transport/authentication signals.
                if (
                    "403" in upper
                    or "429" in upper
                    or "CAPTCHA" in upper
                    or "AUTHENTICATION" in upper
                    or "LOGIN VERIFICATION" in upper
                ):
                    raise

        # ====================================================
        # SAVE ONCE
        # ====================================================

        workbook.save(
            output
        )

    except Exception:

        workbook.close()

        output.unlink(
            missing_ok=True
        )

        raise

    finally:

        try:
            workbook.close()
        except Exception:
            pass

    # ========================================================
    # VERIFY OUTPUT CAN OPEN
    # ========================================================

    verify = load_workbook(
        output,
        data_only=False,
        read_only=True,
    )

    try:

        if (
            sheet_name
            not in verify.sheetnames
        ):
            raise RuntimeError(
                "OUTPUT_VERIFY_SHEET_MISSING:"
                f"{sheet_name}"
            )

    finally:
        verify.close()

    # ========================================================
    # SOURCE INTEGRITY
    # ========================================================

    source_hash_after = (
        _sha256(
            source
        )
    )

    if (
        source_hash_before
        != source_hash_after
    ):
        raise RuntimeError(
            "SOURCE_WORKBOOK_MODIFIED"
        )

    elapsed = (
        time.perf_counter()
        - started
    )

    return ProductWorkbookJobResult(
        source_path=source,
        output_path=output,

        target_count=total,

        written_count=(
            status[
                "WRITTEN"
            ]
        ),

        empty_count=(
            status[
                "EMPTY"
            ]
        ),

        erp_not_found_count=(
            status[
                "ERP_NOT_FOUND"
            ]
        ),

        erp_ambiguous_count=(
            status[
                "ERP_AMBIGUOUS"
            ]
        ),

        error_count=(
            status[
                "ERROR"
            ]
        ),

        written_cell_count=(
            written_cell_count
        ),

        elapsed_seconds=(
            elapsed
        ),

        source_sha256=(
            source_hash_after
        ),

        output_sha256=(
            _sha256(
                output
            )
        ),

        failures=tuple(
            failures
        ),
    )
